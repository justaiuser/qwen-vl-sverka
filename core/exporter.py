# -*- coding: utf-8 -*-
"""
Excel report exporter with professional visual formatting and color-coded discrepancy highlights.
"""
from typing import Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .models import ReconciliationReport, MatchStatus


class ExcelReportExporter:
    def __init__(self):
        # Цветовая палитра
        self.fill_header = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")     # Темно-синий
        self.font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        self.fill_match = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")      # Мягкий зеленый
        self.fill_diff_amt = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")   # Светло-желтый (разница сумм)
        self.fill_missing_1c = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid") # Персиковый (нет у КП)
        self.fill_missing_cp = PatternFill(start_color="EDEDED", end_color="EDEDED", fill_type="solid") # Серый/голубой (нет в 1С)

        self.font_bold = Font(name="Calibri", size=11, bold=True)
        self.font_regular = Font(name="Calibri", size=10)
        self.font_title = Font(name="Calibri", size=14, bold=True, color="1F497D")

        thin = Side(border_style="thin", color="D9D9D9")
        self.border_thin = Border(left=thin, right=thin, top=thin, bottom=thin)

    def export(self, report: ReconciliationReport, output_filepath: str):
        wb = openpyxl.Workbook()
        
        # 1. Сводный отчет
        ws_summary = wb.active
        ws_summary.title = "Сводка"
        self._build_summary_sheet(ws_summary, report)

        # 2. Расхождения
        ws_diff = wb.create_sheet(title="Расхождения")
        self._build_pairs_sheet(ws_diff, [p for p in report.pairs if not p.is_matched], title="ВЕДОМОСТЬ РАСХОЖДЕНИЙ")

        # 3. Сверенные
        ws_matched = wb.create_sheet(title="Сверенные")
        self._build_pairs_sheet(ws_matched, [p for p in report.pairs if p.is_matched], title="СВЕРЕННЫЕ ОПЕРАЦИИ")

        # 4. Все операции
        ws_all = wb.create_sheet(title="Все операции")
        self._build_pairs_sheet(ws_all, report.pairs, title="ПОЛНЫЙ ПРОТОКОЛ СОПОСТАВЛЕНИЯ")

        wb.save(output_filepath)

    def _build_summary_sheet(self, ws, report: ReconciliationReport):
        ws.views.sheetView[0].showGridLines = True

        # Заголовок
        ws.merge_cells("B2:G2")
        ws["B2"] = "ПРОТОКОЛ СВЕРКИ ВЗАИМОРАСЧЕТОВ"
        ws["B2"].font = self.font_title
        ws["B2"].alignment = Alignment(vertical="center")

        # Реквизиты сторон
        ws["B4"] = "Организация (по 1С):"
        ws["C4"] = report.statement_1c.organization or "Наша организация"
        ws["B4"].font = self.font_bold

        ws["B5"] = "Контрагент (по документу):"
        ws["C5"] = report.statement_counterparty.counterparty or report.statement_1c.counterparty or "Контрагент"
        ws["B5"].font = self.font_bold

        ws["B6"] = "Период сверки:"
        ws["C6"] = report.statement_1c.period_str or report.statement_counterparty.period_str or "Не указан"
        ws["B6"].font = self.font_bold

        # Вердикт
        ws.merge_cells("B8:G8")
        ws["B8"] = report.summary_verdict
        ws["B8"].font = Font(name="Calibri", size=12, bold=True, color="006100" if not report.has_discrepancies else "9C0006")
        ws["B8"].fill = PatternFill(start_color="C6EFCE" if not report.has_discrepancies else "FFC7CE", fill_type="solid")
        ws["B8"].alignment = Alignment(horizontal="center", vertical="center")

        # Таблица баланса и оборотов
        headers = ["Показатель", "По данным 1С", "По данным Контрагента", "Расхождение"]
        for col_idx, h in enumerate(headers, start=2):
            cell = ws.cell(row=10, column=col_idx, value=h)
            cell.font = self.font_header
            cell.fill = self.fill_header
            cell.alignment = Alignment(horizontal="center", vertical="center")

        rows_data = [
            ("Сальдо начальное", report.statement_1c.initial_balance, report.statement_counterparty.initial_balance, report.diff_initial_balance),
            ("Обороты Дебет 1С / Кредит КП", report.statement_1c.turnover_debit, report.statement_counterparty.turnover_credit, report.diff_turnover_debit),
            ("Обороты Кредит 1С / Дебет КП", report.statement_1c.turnover_credit, report.statement_counterparty.turnover_debit, report.diff_turnover_credit),
            ("Сальдо конечное", report.statement_1c.final_balance, report.statement_counterparty.final_balance, report.diff_final_balance),
        ]

        for r_offset, (label, v_1c, v_cp, v_diff) in enumerate(rows_data, start=11):
            ws.cell(row=r_offset, column=2, value=label).font = self.font_bold
            
            c_1c = ws.cell(row=r_offset, column=3, value=v_1c)
            c_1c.number_format = '#,##0.00'
            c_1c.font = self.font_regular

            c_cp = ws.cell(row=r_offset, column=4, value=v_cp)
            c_cp.number_format = '#,##0.00'
            c_cp.font = self.font_regular

            c_diff = ws.cell(row=r_offset, column=5, value=v_diff)
            c_diff.number_format = '#,##0.00'
            c_diff.font = self.font_bold
            if abs(v_diff) > 0.01:
                c_diff.font = Font(name="Calibri", size=10, bold=True, color="9C0006")
                c_diff.fill = self.fill_missing_1c

            for col_i in range(2, 6):
                ws.cell(row=r_offset, column=col_i).border = self.border_thin

        # Статистика операций
        ws["B16"] = "СТАТИСТИКА СОПОСТАВЛЕНИЯ"
        ws["B16"].font = self.font_title

        stats = [
            ("Всего операций в 1С:", report.total_items_1c),
            ("Всего операций у контрагента:", report.total_items_cp),
            ("Сверено успешно (совпало):", report.matched_count),
            ("Расхождений (требует внимания):", report.discrepancy_count),
        ]
        for r_offset, (lbl, cnt) in enumerate(stats, start=17):
            ws.cell(row=r_offset, column=2, value=lbl).font = self.font_regular
            c_val = ws.cell(row=r_offset, column=3, value=cnt)
            c_val.font = self.font_bold
            c_val.alignment = Alignment(horizontal="right")

        # Автоширина колонок
        self._autofit_columns(ws)

    def _build_pairs_sheet(self, ws, pairs, title: str):
        ws.views.sheetView[0].showGridLines = True

        # Заголовок таблицы
        ws.merge_cells("A1:K1")
        ws["A1"] = title
        ws["A1"].font = self.font_title
        ws["A1"].alignment = Alignment(vertical="center")

        # Двухуровневая шапка
        ws.merge_cells("A3:B3")
        ws["A3"] = "Результат"
        ws["A3"].font = self.font_header
        ws["A3"].fill = self.fill_header
        ws["A3"].alignment = Alignment(horizontal="center")

        ws.merge_cells("C3:F3")
        ws["C3"] = "Данные 1С:Предприятие"
        ws["C3"].font = self.font_header
        ws["C3"].fill = PatternFill(start_color="244062", fill_type="solid")
        ws["C3"].alignment = Alignment(horizontal="center")

        ws.merge_cells("G3:J3")
        ws["G3"] = "Данные Контрагента (Акт / Скан)"
        ws["G3"].font = self.font_header
        ws["G3"].fill = PatternFill(start_color="366092", fill_type="solid")
        ws["G3"].alignment = Alignment(horizontal="center")

        ws.cell(row=3, column=11, value="Пояснение").font = self.font_header
        ws.cell(row=3, column=11).fill = self.fill_header
        ws.cell(row=3, column=11).alignment = Alignment(horizontal="center")

        subheaders = [
            "Статус", "Разница",
            "Дата 1С", "Документ 1С", "Дебет 1С", "Кредит 1С",
            "Дата КП", "Документ КП", "Дебет КП", "Кредит КП",
            "Примечание"
        ]

        for col_idx, sh in enumerate(subheaders, start=1):
            cell = ws.cell(row=4, column=col_idx, value=sh)
            cell.font = self.font_header
            cell.fill = self.fill_header
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Заполнение строк
        for row_idx, p in enumerate(pairs, start=5):
            # Статус
            status_text = self._status_label(p.status)
            c_stat = ws.cell(row=row_idx, column=1, value=status_text)
            c_stat.alignment = Alignment(horizontal="center")

            # Разница
            c_diff = ws.cell(row=row_idx, column=2, value=p.diff_amount)
            c_diff.number_format = '#,##0.00'

            # 1C
            it_1c = p.item_1c
            ws.cell(row=row_idx, column=3, value=it_1c.date_str if it_1c else "")
            ws.cell(row=row_idx, column=4, value=it_1c.doc_raw if it_1c else "")
            
            c_d1 = ws.cell(row=row_idx, column=5, value=it_1c.debit if it_1c and it_1c.debit else "")
            if it_1c and it_1c.debit: c_d1.number_format = '#,##0.00'
            
            c_c1 = ws.cell(row=row_idx, column=6, value=it_1c.credit if it_1c and it_1c.credit else "")
            if it_1c and it_1c.credit: c_c1.number_format = '#,##0.00'

            # Контрагент
            it_cp = p.item_counterparty
            ws.cell(row=row_idx, column=7, value=it_cp.date_str if it_cp else "")
            ws.cell(row=row_idx, column=8, value=it_cp.doc_raw if it_cp else "")
            
            c_d2 = ws.cell(row=row_idx, column=9, value=it_cp.debit if it_cp and it_cp.debit else "")
            if it_cp and it_cp.debit: c_d2.number_format = '#,##0.00'
            
            c_c2 = ws.cell(row=row_idx, column=10, value=it_cp.credit if it_cp and it_cp.credit else "")
            if it_cp and it_cp.credit: c_c2.number_format = '#,##0.00'

            # Примечание
            ws.cell(row=row_idx, column=11, value=p.explanation)

            # Подсветка строки в зависимости от статуса
            row_fill = self._get_fill_for_status(p.status)
            for c_i in range(1, 12):
                c_cell = ws.cell(row=row_idx, column=c_i)
                c_cell.font = self.font_regular
                c_cell.border = self.border_thin
                if row_fill:
                    c_cell.fill = row_fill

        self._autofit_columns(ws)

    def _status_label(self, status: MatchStatus) -> str:
        mapping = {
            MatchStatus.EXACT: "Сверено",
            MatchStatus.AMOUNT_AND_DOC: "Сдвиг даты",
            MatchStatus.AMOUNT_AND_DATE: "Номер отличен",
            MatchStatus.AMOUNT_DIFF: "Разница сумм",
            MatchStatus.ONLY_1C: "Только в 1С",
            MatchStatus.ONLY_COUNTERPARTY: "Только у КП"
        }
        return mapping.get(status, status.value)

    def _get_fill_for_status(self, status: MatchStatus) -> Optional[PatternFill]:
        if status in (MatchStatus.EXACT, MatchStatus.AMOUNT_AND_DOC, MatchStatus.AMOUNT_AND_DATE):
            return self.fill_match
        if status == MatchStatus.AMOUNT_DIFF:
            return self.fill_diff_amt
        if status == MatchStatus.ONLY_1C:
            return self.fill_missing_1c
        if status == MatchStatus.ONLY_COUNTERPARTY:
            return self.fill_missing_cp
        return None

    def _autofit_columns(self, ws):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len and len(val_str) < 60:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
